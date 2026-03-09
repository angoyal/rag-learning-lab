"""Tests for document readers."""

from __future__ import annotations

from pathlib import Path

import pytest
from src.ingest.readers import (
    read_document,
    read_docx,
    read_html,
    read_markdown,
    read_pdf,
    read_text,
)

# -- Round 1: Text reader --


@pytest.mark.unit
def test_read_text_basic(tmp_path: Path):
    f = tmp_path / "sample.txt"
    f.write_text("Hello, world!")
    doc = read_text(f)
    assert doc.text == "Hello, world!"
    assert doc.source == str(f)
    assert doc.file_type == "txt"
    assert isinstance(doc.metadata, dict)


@pytest.mark.unit
def test_read_text_preserves_whitespace(tmp_path: Path):
    content = "First paragraph.\n\nSecond paragraph.\n  Indented line."
    f = tmp_path / "whitespace.txt"
    f.write_text(content)
    doc = read_text(f)
    assert doc.text == content


@pytest.mark.unit
def test_read_text_empty_file(tmp_path: Path):
    f = tmp_path / "empty.txt"
    f.write_text("")
    doc = read_text(f)
    assert doc.text == ""


@pytest.mark.unit
def test_read_text_file_not_found():
    with pytest.raises(FileNotFoundError):
        read_text(Path("/nonexistent/file.txt"))


# -- Round 2: Markdown reader --


@pytest.mark.unit
def test_read_markdown_strips_formatting(tmp_path: Path):
    f = tmp_path / "test.md"
    f.write_text("This is **bold** and *italic* and [a link](http://example.com).")
    doc = read_markdown(f)
    assert "**" not in doc.text
    assert "*" not in doc.text
    assert "[a link]" not in doc.text
    assert "bold" in doc.text
    assert "italic" in doc.text
    assert "a link" in doc.text
    assert doc.file_type == "markdown"


@pytest.mark.unit
def test_read_markdown_preserves_structure(tmp_path: Path):
    f = tmp_path / "structure.md"
    f.write_text("# Heading\n\nFirst paragraph.\n\nSecond paragraph.")
    doc = read_markdown(f)
    assert "Heading" in doc.text
    assert "First paragraph." in doc.text
    assert "Second paragraph." in doc.text
    # Headings and paragraphs should be separated
    lines = [line for line in doc.text.split("\n") if line.strip()]
    assert len(lines) >= 3


@pytest.mark.unit
def test_read_markdown_with_code_blocks(tmp_path: Path):
    f = tmp_path / "code.md"
    f.write_text("Some text.\n\n```python\nprint('hello')\n```\n\nMore text.")
    doc = read_markdown(f)
    assert "print('hello')" in doc.text


# -- Round 3: HTML reader --


@pytest.mark.unit
def test_read_html_extracts_visible_text(tmp_path: Path):
    f = tmp_path / "test.html"
    f.write_text("<html><body><h1>Title</h1><p>Content here.</p></body></html>")
    doc = read_html(f)
    assert "Title" in doc.text
    assert "Content here." in doc.text
    assert doc.file_type == "html"


@pytest.mark.unit
def test_read_html_strips_scripts_and_styles(tmp_path: Path):
    f = tmp_path / "scripts.html"
    f.write_text(
        "<html><head><style>body{color:red}</style></head>"
        "<body><script>alert('xss')</script><p>Visible text.</p></body></html>"
    )
    doc = read_html(f)
    assert "Visible text." in doc.text
    assert "alert" not in doc.text
    assert "color:red" not in doc.text


@pytest.mark.unit
def test_read_html_handles_nested_tags(tmp_path: Path):
    f = tmp_path / "nested.html"
    f.write_text(
        "<html><body>"
        "<div><div><div><p>Deep <strong>nested</strong> content.</p></div></div></div>"
        "</body></html>"
    )
    doc = read_html(f)
    assert "Deep" in doc.text
    assert "nested" in doc.text
    assert "content." in doc.text


@pytest.mark.unit
def test_read_html_strips_nav_header_footer(tmp_path: Path):
    f = tmp_path / "layout.html"
    f.write_text(
        "<html><body>"
        "<nav>Navigation</nav>"
        "<header>Site Header</header>"
        "<main><p>Main content.</p></main>"
        "<footer>Footer info</footer>"
        "</body></html>"
    )
    doc = read_html(f)
    assert "Main content." in doc.text
    assert "Navigation" not in doc.text
    assert "Site Header" not in doc.text
    assert "Footer info" not in doc.text


# -- Round 4: PDF reader --


@pytest.mark.unit
def test_read_pdf_extracts_text(tmp_path: Path):
    pdf_path = _create_test_pdf(tmp_path / "test.pdf", ["Hello from PDF."])
    doc = read_pdf(pdf_path)
    assert "Hello from PDF" in doc.text
    assert doc.file_type == "pdf"


@pytest.mark.unit
def test_read_pdf_multipage(tmp_path: Path):
    pdf_path = _create_test_pdf(tmp_path / "multi.pdf", ["Page one.", "Page two."])
    doc = read_pdf(pdf_path)
    assert "Page one." in doc.text
    assert "Page two." in doc.text
    # Pages separated by double newline
    assert "\n\n" in doc.text


@pytest.mark.unit
def test_read_pdf_metadata(tmp_path: Path):
    pdf_path = _create_test_pdf(tmp_path / "meta.pdf", ["A", "B", "C"])
    doc = read_pdf(pdf_path)
    assert doc.metadata["page_count"] == 3


# -- Round 5: DOCX reader --


@pytest.mark.unit
def test_read_docx_extracts_paragraphs(tmp_path: Path):
    docx_path = _create_test_docx(tmp_path / "test.docx", ["First paragraph.", "Second paragraph."])
    doc = read_docx(docx_path)
    assert "First paragraph." in doc.text
    assert "Second paragraph." in doc.text
    assert doc.file_type == "docx"


@pytest.mark.unit
def test_read_docx_joins_paragraphs(tmp_path: Path):
    docx_path = _create_test_docx(tmp_path / "join.docx", ["Para one.", "Para two."])
    doc = read_docx(docx_path)
    assert "Para one.\n\nPara two." in doc.text


# -- Round 6: Dispatcher --


@pytest.mark.unit
def test_read_document_routes_txt(tmp_path: Path):
    f = tmp_path / "test.txt"
    f.write_text("plain text")
    doc = read_document(f)
    assert doc.file_type == "txt"
    assert doc.text == "plain text"


@pytest.mark.unit
def test_read_document_routes_md(tmp_path: Path):
    f = tmp_path / "test.md"
    f.write_text("# Title")
    doc = read_document(f)
    assert doc.file_type == "markdown"


@pytest.mark.unit
def test_read_document_routes_html(tmp_path: Path):
    f = tmp_path / "test.html"
    f.write_text("<p>hello</p>")
    doc = read_document(f)
    assert doc.file_type == "html"


@pytest.mark.unit
def test_read_document_routes_htm(tmp_path: Path):
    f = tmp_path / "test.htm"
    f.write_text("<p>hello</p>")
    doc = read_document(f)
    assert doc.file_type == "html"


@pytest.mark.unit
def test_read_document_routes_pdf(tmp_path: Path):
    pdf_path = _create_test_pdf(tmp_path / "test.pdf", ["PDF content."])
    doc = read_document(pdf_path)
    assert doc.file_type == "pdf"


@pytest.mark.unit
def test_read_document_routes_docx(tmp_path: Path):
    docx_path = _create_test_docx(tmp_path / "test.docx", ["DOCX content."])
    doc = read_document(docx_path)
    assert doc.file_type == "docx"


@pytest.mark.unit
def test_read_document_unsupported_extension(tmp_path: Path):
    f = tmp_path / "test.xyz"
    f.write_text("data")
    with pytest.raises(ValueError, match="Unsupported"):
        read_document(f)


@pytest.mark.unit
def test_read_document_case_insensitive(tmp_path: Path):
    f = tmp_path / "test.TXT"
    f.write_text("upper case ext")
    doc = read_document(f)
    assert doc.file_type == "txt"
    assert doc.text == "upper case ext"


# -- Helpers to create test fixtures programmatically --


def _create_test_pdf(path: Path, pages: list[str]) -> Path:
    """Create a minimal PDF with given page texts using fpdf2 or reportlab."""
    try:
        from fpdf import FPDF

        pdf = FPDF()
        for text in pages:
            pdf.add_page()
            pdf.set_font("Helvetica", size=12)
            pdf.cell(text=text)
        pdf.output(str(path))
    except ImportError:
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas

        c = canvas.Canvas(str(path), pagesize=letter)
        for i, text in enumerate(pages):
            if i > 0:
                c.showPage()
            c.drawString(72, 720, text)
        c.save()
    return path


def _create_test_docx(path: Path, paragraphs: list[str]) -> Path:
    """Create a minimal DOCX with given paragraphs."""
    from docx import Document as DocxDocument

    doc = DocxDocument()
    for para in paragraphs:
        doc.add_paragraph(para)
    doc.save(str(path))
    return path
